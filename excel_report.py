from openpyxl import Workbook
from openpyxl.styles import Font


def generate_excel_report(
    results,
    page_summaries
):

    wb = Workbook()

    ws = wb.active

    ws.title = "Changes"

    headers = [
        "Page",
        "Change Type",
        "Severity",
        "Description"
    ]

    for col_num, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=1,
            column=col_num
        )

        cell.value = header
        cell.font = Font(
            bold=True
        )

    row = 2

    for result in results:

        page = result["page"]

        summary = result["summary"]

        severity = "Medium"

        if "removed" in summary.lower():
            severity = "Critical"

        elif "added" in summary.lower():
            severity = "High"

        elif "text integrity" in summary.lower():
            severity = "Low"

        ws.cell(
            row=row,
            column=1
        ).value = page

        ws.cell(
            row=row,
            column=2
        ).value = "AI Detected Change"

        ws.cell(
            row=row,
            column=3
        ).value = severity

        ws.cell(
            row=row,
            column=4
        ).value = summary

        row += 1

    for column in ws.columns:

        max_length = 0

        column_letter = (
            column[0].column_letter
        )

        for cell in column:

            try:

                if len(
                    str(cell.value)
                ) > max_length:

                    max_length = len(
                        str(cell.value)
                    )

            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 5,
            80
        )

    excel_file = (
        "Comparison_Report.xlsx"
    )

    wb.save(
        excel_file
    )

    return excel_file