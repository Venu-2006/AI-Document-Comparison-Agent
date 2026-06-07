from openpyxl import Workbook
from openpyxl.styles import Font


def generate_excel_report_v2(
    page_summaries
):

    wb = Workbook()

    ws = wb.active

    ws.title = "Detailed Changes"

    headers = [

        "Page",

        "Type",

        "Severity",

        "Description"

    ]

    for col_num, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=1,
            column=col_num
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

    row = 2

    for page_num, summary in page_summaries.items():

        lines = summary.split("\n")

        current_section = None

        for line in lines:

            line = line.strip()

            if not line:
                continue

            if "Added Content" in line:

                current_section = "Added Content"
                continue

            elif "Removed Content" in line:

                current_section = "Removed Content"
                continue

            elif "Modified Content" in line:

                current_section = "Modified Content"
                continue

            elif "Layout / Reflow" in line:

                current_section = "Layout / Reflow"
                continue

            elif "Text Integrity Issues" in line:

                current_section = "Text Integrity"
                continue

            elif "Business Impact" in line:

                current_section = None
                continue

            if current_section:

                severity = "Medium"

                if current_section == "Removed Content":
                    severity = "Critical"

                elif current_section == "Added Content":
                    severity = "High"

                elif current_section == "Text Integrity":
                    severity = "Low"

                ws.cell(
                    row=row,
                    column=1
                ).value = page_num

                ws.cell(
                    row=row,
                    column=2
                ).value = current_section

                ws.cell(
                    row=row,
                    column=3
                ).value = severity

                ws.cell(
                    row=row,
                    column=4
                ).value = line

                row += 1

    for column in ws.columns:

        max_length = 0

        column_letter = (
            column[0].column_letter
        )

        for cell in column:

            try:

                if len(
                    str(cell.value)
                ) > max_length:

                    max_length = len(
                        str(cell.value)
                    )

            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 5,
            100
        )

    excel_file = (
        "Comparison_Report_V2.xlsx"
    )

    wb.save(
        excel_file
    )

    return excel_file